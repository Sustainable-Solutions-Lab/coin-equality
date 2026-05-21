"""Create a presentation with the latest figures."""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image


PROJECT_ROOT = Path(__file__).resolve().parent.parent
FIGURES_DIR = PROJECT_ROOT / "data" / "output" / "figures-new"
OUTPUT_PATH = PROJECT_ROOT / "data" / "output" / "figures-new" / "presentation_figures.pptx"

TITLE_COLOR = RGBColor(0x1A, 0x3C, 0x6E)
SUBTITLE_COLOR = RGBColor(0x66, 0x66, 0x66)
TITLE_FONT_SIZE = Pt(28)
SUBTITLE_FONT_SIZE = Pt(18)

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

SLIDES = [
    {
        "title": "How Progressive Taxation and Income-Dependent Damage\nShape the Optimal Carbon Price",
        "subtitle": "May 2026",
        "image": None,
    },
    {
        "title": "Figure 1 — Damage and Tax Distribution",
        "image": FIGURES_DIR / "fig1_damage_tax_distribution_v2.png",
    },
    {
        "title": "Figure 2 — All Cases",
        "image": FIGURES_DIR / "fig2_all_cases_v2.png",
    },
    {
        "title": "Figure 3a — Sensitivity Analysis",
        "image": FIGURES_DIR / "fig3a_sensitivity_4panel_v2.png",
    },
    {
        "title": "Figure 3b — Contour Plots",
        "image": FIGURES_DIR / "fig3b_contour_plots_2panel.png",
    },
    {
        "title": "Figure 4 — Violin Plots",
        "image": FIGURES_DIR / "fig4_violin_plots.png",
    },
]


def add_title_slide(prs, title_text, subtitle_text):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    title_left = Inches(0.8)
    title_top = Inches(2.5)
    title_width = Inches(11.7)
    title_height = Inches(2.0)

    txbox = slide.shapes.add_textbox(title_left, title_top, title_width, title_height)
    tf = txbox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = TITLE_COLOR
    p.alignment = PP_ALIGN.CENTER

    sub_top = Inches(4.8)
    sub_height = Inches(0.6)
    txbox2 = slide.shapes.add_textbox(title_left, sub_top, title_width, sub_height)
    tf2 = txbox2.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = subtitle_text
    p2.font.size = SUBTITLE_FONT_SIZE
    p2.font.color.rgb = SUBTITLE_COLOR
    p2.alignment = PP_ALIGN.CENTER

    return slide


def add_figure_slide(prs, title_text, image_path):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    title_left = Inches(0.8)
    title_top = Inches(0.2)
    title_width = Inches(11.7)
    title_height = Inches(0.7)

    txbox = slide.shapes.add_textbox(title_left, title_top, title_width, title_height)
    tf = txbox.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = TITLE_FONT_SIZE
    p.font.bold = True
    p.font.color.rgb = TITLE_COLOR

    img = Image.open(image_path)
    img_w, img_h = img.size
    aspect = img_w / img_h

    available_width = Inches(11.7)
    available_height = Inches(6.2)

    if available_width / available_height > aspect:
        height = available_height
        width = height * aspect
    else:
        width = available_width
        height = width / aspect

    left = (SLIDE_WIDTH - width) / 2
    top = Inches(1.0) + (available_height - height) / 2

    slide.shapes.add_picture(str(image_path), int(left), int(top), int(width), int(height))

    return slide


def main():
    prs = Presentation()
    prs.slide_width = int(SLIDE_WIDTH)
    prs.slide_height = int(SLIDE_HEIGHT)

    for slide_spec in SLIDES:
        if slide_spec.get("image") is None:
            add_title_slide(prs, slide_spec["title"], slide_spec["subtitle"])
        else:
            add_figure_slide(prs, slide_spec["title"], slide_spec["image"])

    prs.save(str(OUTPUT_PATH))
    print(f"Presentation saved to {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
